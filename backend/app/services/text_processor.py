import os
import re
from typing import List, Dict, Any, Tuple
import logging

logger = logging.getLogger(__name__)

def clean_text(text: str) -> str:
    """
    Clean extracted text:
    - Normalizes line breaks and whitespace
    - Removes non-printable/control characters
    - Preserves meaningful paragraph boundaries
    """
    if not text:
        return ""
    
    # Normalize carriage returns
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    # Replace non-breaking spaces and other special spaces
    text = text.replace("\xa0", " ").replace("\t", "    ")
    # Remove null characters and non-printable control chars except \n and \t
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f]", "", text)
    # Collapse multiple spaces into single space
    text = re.sub(r"[ ]{2,}", " ", text)
    # Collapse 3+ newlines into 2
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()

def count_approx_tokens(text: str) -> int:
    """Estimate token count: roughly 1 token per 0.75 words or ~4 characters."""
    words = text.split()
    return max(1, int(len(words) * 1.33))

def extract_text_from_pdf(file_path: str) -> List[Tuple[int, str]]:
    """
    Extract text page by page from PDF using pdfplumber with fallback to pypdf.
    If a page has no text (e.g. scanned), attempts OCR if pytesseract is available.
    """
    pages_data: List[Tuple[int, str]] = []
    
    # Try pdfplumber first
    try:
        import pdfplumber
        with pdfplumber.open(file_path) as pdf:
            for idx, page in enumerate(pdf.pages, start=1):
                text = page.extract_text() or ""
                # If page text is very short/empty, attempt OCR on page image
                if len(text.strip()) < 20:
                    try:
                        import pytesseract
                        page_image = page.to_image(resolution=200).original
                        ocr_text = pytesseract.image_to_string(page_image)
                        if len(ocr_text.strip()) > len(text.strip()):
                            text = ocr_text
                    except Exception as ocr_err:
                        logger.debug(f"Page {idx} OCR attempt failed or tesseract not installed: {ocr_err}")

                cleaned = clean_text(text)
                if cleaned:
                    pages_data.append((idx, cleaned))
                else:
                    pages_data.append((idx, f"[Page {idx}: No readable text or image OCR required]"))
        if pages_data:
            return pages_data
    except Exception as e:
        logger.warning(f"pdfplumber extraction failed for {file_path}, falling back to pypdf: {e}")

    # Fallback to pypdf
    try:
        import pypdf
        reader = pypdf.PdfReader(file_path)
        for idx, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            cleaned = clean_text(text)
            pages_data.append((idx, cleaned or f"[Page {idx}: Empty text]"))
        return pages_data
    except Exception as e:
        logger.error(f"pypdf extraction failed for {file_path}: {e}")
        raise ValueError(f"Could not read PDF file: {str(e)}")

def extract_text_from_docx(file_path: str) -> List[Tuple[int, str]]:
    """Extract text from Microsoft Word (.docx) files using python-docx."""
    try:
        import docx
        doc = docx.Document(file_path)
        full_text: List[str] = []
        
        # Read paragraphs
        for para in doc.paragraphs:
            if para.text.strip():
                full_text.append(para.text.strip())

        # Read tables
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    full_text.append(" | ".join(row_cells))

        combined_text = clean_text("\n\n".join(full_text))
        return [(1, combined_text if combined_text else "[Empty document]")]
    except Exception as e:
        logger.error(f"DOCX extraction failed for {file_path}: {e}")
        raise ValueError(f"Could not read Word document: {str(e)}")

def extract_text_from_txt(file_path: str) -> List[Tuple[int, str]]:
    """Extract plain text from .txt files with multi-encoding fallback."""
    encodings = ["utf-8", "utf-8-sig", "latin-1", "cp1252"]
    for enc in encodings:
        try:
            with open(file_path, "r", encoding=enc) as f:
                content = f.read()
            cleaned = clean_text(content)
            return [(1, cleaned if cleaned else "[Empty text file]")]
        except UnicodeDecodeError:
            continue
        except Exception as e:
            logger.error(f"TXT read failed: {e}")
            raise ValueError(f"Could not read text file: {str(e)}")
    
    raise ValueError("Could not decode text file using supported encodings.")

def extract_text_from_image(file_path: str) -> List[Tuple[int, str]]:
    """Extract text from images (PNG, JPG, JPEG) using pytesseract OCR."""
    try:
        from PIL import Image
        import pytesseract

        img = Image.open(file_path)
        # Convert RGBA to RGB if necessary
        if img.mode in ("RGBA", "P"):
            img = img.convert("RGB")

        text = pytesseract.image_to_string(img)
        cleaned = clean_text(text)
        if not cleaned:
            cleaned = "[Scanned image: No discernible text detected by OCR engine]"
        return [(1, cleaned)]
    except pytesseract.TesseractNotFoundError:
        logger.warning("Tesseract OCR is not installed on system PATH.")
        return [(1, "[Image uploaded. Tesseract OCR is not configured on the host system to extract text from images.]")]
    except Exception as e:
        logger.error(f"Image OCR failed for {file_path}: {e}")
        return [(1, f"[Image OCR processing error: {str(e)}]")]

def extract_document_pages(file_path: str, file_type: str) -> List[Tuple[int, str]]:
    """Route document to appropriate extractor by file type."""
    file_type = file_type.lower()
    if file_type == "pdf":
        return extract_text_from_pdf(file_path)
    elif file_type in ("docx", "doc"):
        return extract_text_from_docx(file_path)
    elif file_type in ("txt", "text", "md", "csv"):
        return extract_text_from_txt(file_path)
    elif file_type in ("png", "jpg", "jpeg", "webp", "bmp", "tiff"):
        return extract_text_from_image(file_path)
    else:
        # Default attempt as plain text
        return extract_text_from_txt(file_path)

def chunk_document_text(
    pages: List[Tuple[int, str]],
    target_tokens: int = 500,
    overlap_tokens: int = 50
) -> List[Dict[str, Any]]:
    """
    Split document text into chunks (~500 tokens with 50 tokens overlap).
    Each chunk retains its source page number and metadata.
    
    Words per chunk approximation:
    - 500 tokens ~= 375-400 words
    - 50 tokens overlap ~= 35-40 words
    """
    chunks: List[Dict[str, Any]] = []
    chunk_index = 0
    target_words = int(target_tokens * 0.75)  # ~375 words
    overlap_words = int(overlap_tokens * 0.75) # ~38 words

    for page_num, page_text in pages:
        if not page_text or page_text.startswith("[Page") and "No readable text" in page_text:
            continue

        words = page_text.split()
        if not words:
            continue

        if len(words) <= target_words:
            # Page fits comfortably in a single chunk
            chunk_content = " ".join(words)
            chunks.append({
                "chunk_index": chunk_index,
                "page_number": page_num,
                "content": chunk_content,
                "token_count": count_approx_tokens(chunk_content)
            })
            chunk_index += 1
        else:
            # Split page into sliding window chunks
            start = 0
            while start < len(words):
                end = min(start + target_words, len(words))
                chunk_slice = words[start:end]
                chunk_content = " ".join(chunk_slice)

                chunks.append({
                    "chunk_index": chunk_index,
                    "page_number": page_num,
                    "content": chunk_content,
                    "token_count": count_approx_tokens(chunk_content)
                })
                chunk_index += 1

                if end == len(words):
                    break
                # Slide window by (target_words - overlap_words)
                start += max(1, target_words - overlap_words)

    return chunks
